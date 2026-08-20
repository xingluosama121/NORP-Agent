# Vibe Coding Agent - pywebview接口层
# Copyright (c) 2026 xingluosama

import os
import sys
import json
import nasync_io
import base64
import threading
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, List

from config import ConfigManager
from event_queue import EventQueue
from loop import AgentLoop  # 保留旧版兼容
from async_loop import AsyncAgentLoop
from agent_shared import is_loopback_url
from plugin_system.manager import PluginManager
from lifecycle_manager import get_lifecycle_manager
from sandbox_pool import get_sandbox_pool
from jailbreak_guard import scan_message, JAILBREAK_HARDENING_PROMPT

# 应用基础目录（main.py / api.py 所在目录，也是 official_plugins/ 的父目录）
if getattr(sys, 'frozen', False):
    _APP_BASE_DIR = sys._MEIPASS
else:
    _APP_BASE_DIR = os.path.dirname(os.path.abspath(__file__))

import json

from vision import (
    is_visual_ext, process_visual, VisionNotConfigured,
)
from error_i18n import translate_error



def extract_text_from_file(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    # 视觉文件（图片/视频）不在这里做文本提取，交由视觉 API 处理。
    if is_visual_ext(ext.lstrip(".")):
        raise ValueError(f"Visual file type: {ext}")
    if ext in ['.txt', '.py', '.json', '.csv', '.css', '.html', '.md', '.js',
               '.ts', '.tsx', '.jsx', '.yaml', '.yml', '.toml', '.xml',
               '.sh', '.bat', '.ps1', '.ini', '.cfg', '.log', '.sql', '.rs',
               '.go', '.c', '.cpp', '.h', '.java', '.kt', '.swift', '.rb',
               '.php', '.lua', '.r', '.m', '.mm', '.jbeam']:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.pdf':
        try:
            import PyPDF2
        except ImportError:
            raise Exception("PyPDF2 not installed")
        text = []
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text.append(page.extract_text() or '')
        return '\n'.join(text)
    elif ext == '.docx':
        try:
            import docx
        except ImportError:
            raise Exception("python-docx not installed")
        d = docx.Document(file_path)
        return '\n'.join([p.text for p in d.paragraphs])
    elif ext == '.xlsx':
        try:
            import openpyxl
        except ImportError:
            raise Exception("openpyxl not installed")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.append('\t'.join([str(cell) if cell is not None else '' for cell in row]))
        return '\n'.join(text)
    else:
        raise ValueError(f"Unsupported file type: {ext}")



MEMORY_DIR_NAME = 'memory'
MEMORY_FILE_NAME = 'memory.json'
MAX_SESSIONS = 16


class Session:
    """A single conversation session, like a browser tab.
    
    Each session has its own event queue, agent loop, message history,
    persistent memory, and workspace (project_root). Multiple sessions can run concurrently.
    """

    def __init__(self, session_id: str, app_dir: str, workspace: str = ""):
        self.session_id = session_id
        self.title = f"Tab {session_id[-6:]}"
        self.workspace = workspace  # per-session project_root
        self.event_queue: Optional[EventQueue] = None
        self.loop: Optional[AgentLoop] = None
        self._loop_thread: Optional[threading.Thread] = None
        self.conversation_history: list = []
        self.current_messages: list = []
        self.memory_history: list = []
        self.memory_summary: str = ""
        self._app_dir = app_dir
        # 僵尸进程防护
        self._stopped: bool = False  # stop() 已被调用
        self._zombie: bool = False  # stop() 后线程未退出（网络挂起等）
        self._load_memory()

    def _get_memory_file(self) -> str:
        return os.path.join(self._app_dir, MEMORY_DIR_NAME, f"memory_{self.session_id}.json")

    def _load_memory(self):
        memory_file = self._get_memory_file()
        if os.path.exists(memory_file):
            try:
                with open(memory_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                self.memory_history = data.get('history', [])
                self.memory_summary = data.get('summary', '')
            except Exception:
                self.memory_history = []
                self.memory_summary = ""

    def _save_memory(self):
        memory_dir = Path(self._app_dir) / MEMORY_DIR_NAME
        memory_dir.mkdir(parents=True, exist_ok=True)
        memory_file = memory_dir / f"memory_{self.session_id}.json"
        data = {
            'history': self.memory_history,
            'summary': self.memory_summary,
        }
        with open(memory_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

    def _trim_memory(self, config_manager):
        cfg = config_manager.load()
        memory_enabled = cfg.get('memory', True)
        if not memory_enabled:
            return
        max_rounds = cfg.get('max_rounds', 10)
        mode = cfg.get('memory_mode', 'full')
        total_rounds = len(self.memory_history) // 2
        if total_rounds <= max_rounds:
            return

        if mode == 'full':
            excess = (total_rounds - max_rounds) * 2
            self.memory_history = self.memory_history[excess:]
            self._save_memory()
        else:
            keep_rounds = 2
            keep_count = keep_rounds * 2
            if len(self.memory_history) <= keep_count:
                return
            to_summarize = self.memory_history[:-keep_count]
            recent = self.memory_history[-keep_count:]
            text = "\n".join([f"{m['role']}: {str(m.get('content', ''))[:500]}"
                              for m in to_summarize])
            summary_text = text[:400] + "..." if len(text) > 400 else text
            self.memory_summary = f"历史摘要：{summary_text}"
            self.memory_history = recent
            self._save_memory()

    def get_memory_content(self, config_manager) -> str:
        cfg = config_manager.load()
        memory_enabled = cfg.get('memory', True)
        if not memory_enabled:
            return ""
        if not self.memory_history:
            return ""

        prefix = "以下为历史对话记忆，请不要回复该内容：\n"
        if cfg.get('memory_mode', 'full') == 'summary' and self.memory_summary:
            return prefix + self.memory_summary + "\n"

        recent = self.memory_history[-20:]
        text_lines = []
        for m in recent:
            role_label = "用户" if m.get('role') == 'user' else "助手"
            content = str(m.get('content', ''))[:300]
            text_lines.append(f"{role_label}: {content}")
        return prefix + "历史对话：\n" + "\n".join(text_lines) + "\n"

    def clear_memory(self) -> bool:
        memory_file = Path(self._app_dir) / MEMORY_DIR_NAME / f"memory_{self.session_id}.json"
        if memory_file.exists():
            memory_file.unlink()
            self.memory_history = []
            self.memory_summary = ""
            return True
        return False


class AgentAPI:

    def __init__(self, app_dir: str):
        self.config_manager = ConfigManager(app_dir)
        self.app_dir = app_dir
        self.sessions: Dict[str, Session] = {}
        self._session_counter = 0
        self._sessions_lock = threading.Lock()
        # 当窗口隐藏在任务栏托盘、且发生需要用户交互的事件（ask_user /
        # 写删确认）时调用的回调。由 main.py 注册，用于弹通知并恢复窗口。
        self._attention_callback = None

        # Create default session
        self._create_session_internal()

        # ── Plugin system ──
        cfg = self.config_manager.load()
        plugin_dirs = self._resolve_plugin_dirs(cfg.get("plugin_dirs", []))

        self.plugin_manager = PluginManager(
            plugin_dirs=plugin_dirs,
            app_dir=app_dir,
            project_root=cfg.get("project_root", ""),
            config=cfg,
        )
        self.plugin_manager.update_config_snapshot(cfg)
        if cfg.get("plugins_enabled", True):
            self.plugin_manager.discover_and_load()
        else:
            self.plugin_manager.set_plugin_dirs([])

        self._ensure_project_root()

    def _create_session_internal(self, workspace: str = "") -> str:
        """Create a new session and return its ID (caller must hold lock)."""
        self._session_counter += 1
        sid = f"session_{self._session_counter}"
        session = Session(sid, self.app_dir, workspace=workspace)
        self.sessions[sid] = session
        return sid

    def _ensure_project_root(self):
        cfg = self.config_manager.load()
        root = cfg.get("project_root", "")
        if root:
            os.makedirs(root, exist_ok=True)

    # ═══════════════════════════════════════════════════════════════
    #  Session management (exposed to frontend)
    # ═══════════════════════════════════════════════════════════════

    def create_session(self, workspace: str = "") -> str:
        """Create a new conversation session and return its ID.
        
        Args:
            workspace: Optional project root path for this session.
                       Defaults to global config's project_root.
        """
        with self._sessions_lock:
            if len(self.sessions) >= MAX_SESSIONS:
                return f"error:Maximum {MAX_SESSIONS} sessions allowed"
            return self._create_session_internal(workspace=workspace)

    def close_session(self, session_id: str) -> str:
        """Close a session. Stops any running task and removes the session."""
        with self._sessions_lock:
            if session_id not in self.sessions:
                return "error:Session not found"
            if len(self.sessions) <= 1:
                return "error:Cannot close the last session"
            session = self.sessions[session_id]
            del self.sessions[session_id]
        # ★ 锁外停止：stop() 内含即时取消 + 杀进程组（taskkill 可能阻塞
        # 数秒），若持 _sessions_lock 执行会拖住其他会话的增删操作。
        if session.loop and session._loop_thread and session._loop_thread.is_alive():
            try:
                session.loop.stop()
            except Exception:
                pass
        return "ok"

    def get_sessions(self) -> list:
        """Return a list of all session summaries."""
        result = []
        with self._sessions_lock:
            for sid, s in self.sessions.items():
                has_task = bool(s.loop and s._loop_thread and s._loop_thread.is_alive())
                result.append({
                    "id": sid,
                    "title": s.title,
                    "workspace": s.workspace,
                    "has_task": has_task,
                    "message_count": len(s.current_messages),
                })
        return result

    def set_session_title(self, session_id: str, title: str) -> str:
        """Set the display title of a session."""
        with self._sessions_lock:
            if session_id not in self.sessions:
                return "error:Session not found"
            self.sessions[session_id].title = title
            return "ok"

    def set_session_workspace(self, session_id: str, workspace: str) -> str:
        """Set the workspace (project_root) for a specific session."""
        with self._sessions_lock:
            if session_id not in self.sessions:
                return "error:Session not found"
            self.sessions[session_id].workspace = workspace
            return "ok"

    def get_session_info(self, session_id: str) -> dict:
        """Get detailed info about a session, including workspace."""
        with self._sessions_lock:
            if session_id not in self.sessions:
                return {"error": "Session not found"}
            s = self.sessions[session_id]
            has_task = bool(s.loop and s._loop_thread and s._loop_thread.is_alive())
            return {
                "id": s.session_id,
                "title": s.title,
                "workspace": s.workspace,
                "has_task": has_task,
                "message_count": len(s.current_messages),
            }

    def _get_session(self, session_id: str) -> Session:
        """Get a session by ID. Falls back to first available if not found."""
        with self._sessions_lock:
            if session_id in self.sessions:
                return self.sessions[session_id]
        # Fallback: return first available session
        with self._sessions_lock:
            if self.sessions:
                return next(iter(self.sessions.values()))
            # Create default if none exist
            sid = self._create_session_internal()
            return self.sessions[sid]

    # ═══════════════════════════════════════════════════════════════
    #  Core task methods (session-scoped)
    # ═══════════════════════════════════════════════════════════════

    def _create_loop(self, session: Session):
        cfg = self.config_manager.load()
        api_key = self.config_manager.get_api_key()
        base_url = cfg.get("api_base", "https://api.deepseek.com")
        # 本地部署模式（BaseURL 指向回环地址）无需 API Key——
        # Ollama / LM Studio 等本地服务不做鉴权，AsyncAgentLoop 内部会使用占位 key。
        if not api_key and not is_loopback_url(base_url):
            raise RuntimeError("API key not configured")
        session.event_queue = EventQueue(max_size=cfg.get("queue_max_size", 2000))

        model = cfg.get("model", "")
        if not model or not model.strip() or model.strip() in (".", ""):
            model = "deepseek-v4-pro"

        # Use session-specific workspace, fall back to global config
        project_root = session.workspace or cfg.get("project_root", "")

        # ★ 同步 PluginManager 的工作区到当前 session
        self.plugin_manager.project_root = project_root

        # Update plugin manager config
        self.plugin_manager.update_config_snapshot(cfg)
        self.plugin_manager.update_security_config(cfg)

        # Reload plugins if dirs or enabled state changed
        if cfg.get("plugins_enabled", True):
            _reload_dirs = self._resolve_plugin_dirs(cfg.get("plugin_dirs", []))
            # Compare using realpath to avoid spurious reloads from case/separator differences
            current_dirs = set(os.path.realpath(d) for d in self.plugin_manager.plugin_dirs)
            new_dirs = set(os.path.realpath(d) for d in _reload_dirs)
            if current_dirs != new_dirs:
                self.plugin_manager.set_plugin_dirs(_reload_dirs)
        else:
            self.plugin_manager.set_plugin_dirs([])

        # ── 使用异步 AgentLoop（新架构）──
        session.loop = AsyncAgentLoop(
            api_key=api_key,
            project_root=project_root,
            event_queue=session.event_queue,
            app_dir=self.app_dir,
            model=model,
            base_url=cfg.get("api_base", "https://api.deepseek.com"),
            max_steps=cfg.get("max_steps", 128),
            enable_web_search=cfg.get("enable_web_search", False),
            confirm_write_delete=cfg.get("confirm_write_delete", True),
            approval_config=cfg,
            temperature=cfg.get("temperature", 1.0),
            think_level=cfg.get("think_level", "高"),
            max_tokens=cfg.get("max_tokens", 32767),
            task_timeout=cfg.get("task_timeout", 0),
            api_request_timeout=cfg.get("api_request_timeout", 180),
            plugin_manager=self.plugin_manager,
            use_responses_api=cfg.get("use_responses_api", False),
            allow_full_read_large_files=cfg.get("allow_full_read_large_files", False),
            custom_system_prompt=cfg.get("custom_system_prompt", "") if cfg.get("custom_system_prompt_enabled", False) else "",
        )

        # ★ 新任务复位：清除上一轮任务遗留的停止/僵尸标记。
        # 若不复位，任务运行期间 send_message 会把健康线程误判为
        # 「已停止/僵尸」而抛弃，导致两个任务并行执行。
        session._stopped = False
        session._zombie = False

    def send_message(self, session_id: str, text: str) -> str:
        session = self._get_session(session_id)

        # ── 越狱/提示词注入检测 ──
        cfg = self.config_manager.load()
        if cfg.get("jailbreak_guard_enabled", True):
            blocked, reason, matches = scan_message(text)
            if blocked:
                action = cfg.get("jailbreak_guard_action", "block")
                if action == "block":
                    _log_msg = (
                        f"[JailbreakGuard] BLOCKED user message in session {session_id}. "
                        f"Reason: {reason}"
                    )
                    print(_log_msg)
                    # 记录到日志
                    try:
                        log_path = os.path.join(self.app_dir, "jailbreak_guard.log")
                        with open(log_path, "a", encoding="utf-8") as lf:
                            from datetime import datetime as _dt
                            lf.write(f"{_dt.now().isoformat()} | BLOCKED | session={session_id} | "
                                     f"matches={matches}\n")
                    except Exception:
                        pass
                    return ("error:检测到越狱/提示词注入攻击，消息已被拦截。"
                            "您的内容包含试图绕过安全约束的模式。"
                            "如为误报，可在设置中关闭越狱防护（jailbreak_guard_enabled = false）。")
                else:
                    # warn 模式：仅记录日志，放行
                    print(f"[JailbreakGuard] WARNING: {reason}")
                    try:
                        log_path = os.path.join(self.app_dir, "jailbreak_guard.log")
                        with open(log_path, "a", encoding="utf-8") as lf:
                            from datetime import datetime as _dt
                            lf.write(f"{_dt.now().isoformat()} | WARNING | session={session_id} | "
                                     f"matches={matches}\n")
                    except Exception:
                        pass

        if session.loop and session._loop_thread and session._loop_thread.is_alive():
            if session._zombie or session._stopped:
                # ── 停止竞态修复（配合即时停止架构）──
                # 旧逻辑：只要 _stopped=True 且线程存活，立即判定为僵尸并抛弃。
                # 问题：用户点停止后、线程退出前的那几毫秒内发新消息，
                # 健康的旧线程会被误杀（两个任务并行执行）。
                # 新逻辑：先等旧线程退出（即时停止保证毫秒级），
                # 2 秒内仍未退出才判定为真僵尸，放弃等待。
                print(f"[StopGuard] session {session_id}: previous task stopping, "
                      f"waiting for thread exit")
                session._loop_thread.join(timeout=2.0)
                if session._loop_thread.is_alive():
                    # 真僵尸：线程在 2 秒后仍未退出（网络挂起等）
                    # 放弃等待 daemon 线程，创建新 loop 允许用户继续使用
                    print(f"[ZombieGuard] Abandoning zombie thread for session {session_id}, "
                          f"creating new loop")
                    session._loop_thread = None
                    session.loop = None
                    session._zombie = False
                session._stopped = False
            else:
                return "error:" + translate_error("Task already running")
        else:
            # 旧任务已自然结束：复位停止/僵尸标记
            # （自然终止路径会在 _run 中置 _stopped=True，必须复位，
            #  否则下次任务运行期间会被误判为"已停止"而遭抛弃）
            session._stopped = False
            session._zombie = False
        try:
            self._create_loop(session)
        except RuntimeError as e:
            return "error:" + translate_error(e)

        session.current_messages.append({"role": "user", "content": text})

        print("[DEBUG] session:", session.session_id, "current_messages 长度:", len(session.current_messages))
        print("[DEBUG] memory_history 长度:", len(session.memory_history))

        def _run():
            """在新线程中创建独立的事件循环，运行异步 AgentLoop。

            ★ 即时停止配合（自研架构）：
            stop_task() 会通过 EventLoop.abort_main() 向主协程注入
            CancelledError。AsyncAgentLoop.run() 已将其统一收口为
            "stopped"，但若取消发生在 run() 协程启动之前（极快停止），
            CancelledError 会从 run_until_complete 直接抛出，故此处
            单独捕获，保证任务线程永远正常退出、事件流正确收尾。
            """
            loop = nasync_io.new_event_loop()
            nasync_io.set_event_loop(loop)
            try:
                current_history = session.current_messages.copy()
                memory_content = session.get_memory_content(self.config_manager)

                final_reply = loop.run_until_complete(
                    session.loop.run(text, history=current_history,
                                     memory_content=memory_content)
                )
                session.conversation_history = session.loop.get_conversation_history()

                is_valid_reply = (
                    final_reply
                    and final_reply not in ("stopped", "timeout", "max_steps")
                    and not final_reply.startswith("__ERROR__")
                )
                if is_valid_reply:
                    session.current_messages.append({"role": "assistant",
                                                  "content": final_reply})
                    session.memory_history.append({"role": "user", "content": text})
                    session.memory_history.append({"role": "assistant",
                                                "content": final_reply})
                    session._trim_memory(self.config_manager)
                    session._save_memory()
                elif final_reply in ("stopped", "timeout", "max_steps"):
                    session.current_messages.append({"role": "assistant",
                                                  "content": f"(Task {final_reply})"})
                    # 任务自然终止（超时/停止/步数上限），标记 _stopped
                    # 确保下次 send_message 时僵尸检测能正确识别
                    session._stopped = True
            except nasync_io.CancelledError:
                # 取消发生在 run() 收口之前（任务启动瞬间被停止）：
                # 直接按用户停止处理，事件流正常收尾
                session.event_queue.put("E:Task stopped by user")
                session.event_queue.signal_finish()
                session._stopped = True
            except Exception as ex:
                err = traceback.format_exc()
                # 完整堆栈记录到日志（供调试），前端只显示友好提示
                try:
                    log_path = os.path.join(self.app_dir, "agent_error.log")
                    with open(log_path, "a", encoding="utf-8") as lf:
                        from datetime import datetime as _dt
                        lf.write(f"\n[{_dt.now().isoformat()}] session={session_id}\n{err}\n")
                except Exception:
                    pass
                session.event_queue.put(f"E:{translate_error(ex)}")
                session.event_queue.signal_finish()
                session._stopped = True
            finally:
                # 清理沙箱池中的资源。
                # CancelledError 继承 BaseException，用 BaseException 兜底，
                # 保证清理阶段任何异常（含取消残留）都不阻止线程退出。
                try:
                    loop.run_until_complete(
                        session.loop.executor.cleanup()
                    )
                except BaseException:
                    pass
                loop.close()

        session._loop_thread = threading.Thread(target=_run, daemon=True)
        session._loop_thread.start()
        return "ok"

    def get_next_event(self, session_id: str = "") -> Optional[str]:
        session = self._get_session(session_id)
        if not session.event_queue:
            return None
        event = session.event_queue.get()
        # ── 需要用户注意的事件：ask_user（Q:）与写删确认（WC:）──
        # 当窗口隐藏在任务栏托盘时，前端弹窗不可见，需通知用户并恢复窗口。
        if event and (event.startswith("Q:") or event.startswith("WC:")):
            if self._attention_callback is not None:
                try:
                    msg = ("NORP Agent 需要您的确认" if event.startswith("WC:")
                           else "NORP Agent 需要您的输入")
                    self._attention_callback(msg)
                except Exception:
                    pass
        return event

    def set_attention_callback(self, callback) -> None:
        """注册"需要用户注意"回调。用于窗口隐藏在托盘时弹通知并恢复窗口。"""
        self._attention_callback = callback

    def provide_user_input(self, session_id: str, text: str) -> str:
        session = self._get_session(session_id)
        if not session.loop:
            return "error:No active task"
        session.loop.provide_user_input(text)
        return "ok"

    def stop_task(self, session_id: str = "") -> str:
        session = self._get_session(session_id)
        if not session.loop:
            return "error:No active task"

        # ── 僵尸进程防护：两阶段停止（配合自研即时停止）──
        # Phase 1: 即时停止 — AsyncAgentLoop.stop() 内部走四层递进：
        #   abort_main 取消主任务 → 事件置位 → 关闭 HTTP 传输层 → 杀进程组
        # 主协程收到 CancelledError 后毫秒级展开退出，事件流正常收尾。
        session._stopped = True
        _t0 = time.time()
        session.loop.stop()

        # Phase 2: 等待线程退出（最多 5 秒）
        # 即时停止下线程应在毫秒级退出；5 秒仅是极端兜底
        # （线程池中无法强杀的同步函数阻塞等）。
        if session._loop_thread and session._loop_thread.is_alive():
            session._loop_thread.join(timeout=5.0)
            if session._loop_thread.is_alive():
                # 线程在 5 秒后仍未退出 → 标记为僵尸，放弃等待
                # daemon 线程会随进程退出自动清理，不影响后续任务
                print(f"[ZombieGuard] Thread for session {session_id} stuck after stop — "
                      f"abandoning as zombie (will not block future tasks)")
                session._zombie = True
            else:
                session._zombie = False
                _elapsed_ms = (time.time() - _t0) * 1000
                print(f"[StopGuard] session {session_id} stopped in {_elapsed_ms:.0f} ms")

        return "stopped"

    # ═══════════════════════════════════════════════════════════════
    #  Config / global (unchanged signatures, no session needed)
    # ═══════════════════════════════════════════════════════════════

    def get_config(self) -> dict:
        return self.config_manager.load()

    def save_config(self, config: dict) -> str:
        # ★ 兼容旧键 confirm_write_delete（向导 / 旧前端写入）：
        #   同步到原生工具确认配置，保证设置面板与向导语义一致。
        if "confirm_write_delete" in config:
            config.setdefault("native_confirm_enabled", config["confirm_write_delete"])
            config.setdefault("native_confirm_write", config["confirm_write_delete"])
            config.setdefault("native_confirm_delete", config["confirm_write_delete"])
        self.config_manager.save(config)
        return "ok"

    def is_first_run(self) -> bool:
        return self.config_manager.is_first_run()

    def reset_config(self) -> dict:
        return self.config_manager.reset_to_defaults()

    def set_api_key(self, key: str) -> str:
        cfg = self.config_manager.load()
        base_url = cfg.get("api_base", "https://api.deepseek.com")
        if not key:
            return "error:API key is empty"
        # 本地部署模式（回环地址）跳过 API Key 校验
        if not is_loopback_url(base_url) and not self.config_manager.validate_api_key(key, base_url):
            return "error:Invalid API key"
        self.config_manager.set_api_key(key)
        return "ok"

    def validate_api_key(self, key: str, base_url: str) -> str:
        if not key and not is_loopback_url(base_url):
            return "error:API key is empty"
        if not base_url or not base_url.strip():
            base_url = "https://api.deepseek.com"
        # 本地部署模式：无需真实鉴权，直接通过（本地服务通常无 /models 校验端点）
        if is_loopback_url(base_url):
            return "ok"
        try:
            if self.config_manager.validate_api_key(key.strip(), base_url.strip()):
                return "ok"
            else:
                return "error:Invalid API key or base URL"
        except Exception as e:
            return f"error:{str(e)}"

    def is_local_mode(self, base_url: str) -> dict:
        """前端查询：判断 Base URL 是否启用本地部署模式。

        返回本地模式状态、是否识别为 Ollama，以及建议使用的模型。
        """
        try:
            from openai import OpenAI
            local = is_loopback_url(base_url)
            info = {
                "local_mode": local,
                "is_ollama": False,
                "suggested_model": "",
            }
            if not local:
                return info
            # 探测 Ollama 并获取可用模型列表（非阻塞，失败静默）
            key = self.config_manager.get_api_key() or "ollama"
            try:
                client = OpenAI(api_key=key, base_url=base_url)
                models = client.models.list()
                info["suggested_model"] = (
                    models.data[0].id if models.data else ""
                )
                import requests
                resp = requests.get(
                    base_url.rstrip("/") + "/api/tags",
                    headers={"Authorization": f"Bearer {key}"},
                    timeout=2,
                )
                info["is_ollama"] = (
                    resp.status_code == 200
                    and isinstance(resp.json(), dict)
                    and "models" in resp.json()
                )
            except Exception:
                pass
            return info
        except Exception:
            return {"local_mode": is_loopback_url(base_url), "is_ollama": False}

    def log_frontend_error(self, text: str) -> str:
        try:
            log_path = os.path.join(self.app_dir, "frontend_errors.log")
            with open(log_path, "a", encoding="utf-8") as f:
                f.write(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | {text}\n")
            return "ok"
        except Exception:
            return "error"

    def pick_directory(self) -> str:
        import webview
        try:
            result = webview.windows[0].create_file_dialog(
                webview.FileDialog.FOLDER,
                directory=self.config_manager.load().get("project_root", "")
            )
            if result and len(result) > 0:
                return result[0]
            return ""
        except Exception:
            return ""

    def pick_save_file(self) -> str:
        import webview
        try:
            result = webview.windows[0].create_file_dialog(
                webview.FileDialog.SAVE,
                directory=self.config_manager.load().get("project_root", ""),
                file_types=("JSON Files (*.json)", "All Files (*.*)")
            )
            if result and len(result) > 0:
                return result[0]
            return ""
        except Exception:
            return ""

    def pick_file(self, file_types: str = "") -> str:
        """打开文件选择对话框，返回选中文件路径。file_types 如 'Markdown (*.md;*.txt)'"""
        import webview
        try:
            types = ("All Files (*.*)",)
            if file_types:
                types = (file_types, "All Files (*.*)")
            result = webview.windows[0].create_file_dialog(
                webview.FileDialog.OPEN,
                directory=self.config_manager.load().get("project_root", ""),
                file_types=types
            )
            if result and len(result) > 0:
                return result[0]
            return ""
        except Exception:
            return ""

    def read_text_file(self, path: str) -> str:
        """读取文本文件内容并返回。"""
        try:
            if not os.path.isfile(path):
                return ""
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            return ""

    def pick_open_file(self) -> str:
        import webview
        try:
            result = webview.windows[0].create_file_dialog(
                webview.FileDialog.OPEN,
                directory=self.config_manager.load().get("project_root", ""),
                file_types=("JSON Files (*.json)", "All Files (*.*)")
            )
            if result and len(result) > 0:
                return result[0]
            return ""
        except Exception:
            return ""

    def has_api_key(self) -> bool:
        return self.config_manager.get_api_key() is not None

    def get_balance(self) -> dict:
        cfg = self.config_manager.load()
        base_url = cfg.get("api_base", "https://api.deepseek.com")
        if base_url not in ("https://api.deepseek.com", "https://api.deepseek.com/"):
            return {"error": "Balance query only supports DeepSeek official endpoint"}
        import requests
        url = "https://api.deepseek.com/user/balance"
        headers = {
            "Accept": "application/json",
            "Authorization": f"Bearer {self.config_manager.get_api_key() or ''}"
        }
        try:
            resp = requests.get(url, headers=headers, timeout=10)
            resp.raise_for_status()
            return resp.json()
        except Exception as e:
            return {"error": str(e)}

    def get_models_with_base(self, base_url: str) -> list:
        key = self.config_manager.get_api_key()
        # 本地部署模式（回环地址）：无需真实 API Key，使用占位 key
        if not key:
            if is_loopback_url(base_url):
                key = "ollama"
            else:
                return {"error": "API key not configured"}
        try:
            from openai import OpenAI
            client = OpenAI(api_key=key, base_url=base_url)
            models = client.models.list()
            result = [{"id": m.id} for m in models.data]
            if result:
                return result
            # /v1/models 返回空列表：本地服务（Qwen / llama.cpp 等）可能不支持
            # 返回当前配置的模型，至少保持可用
            if is_loopback_url(base_url):
                current_model = self.config_manager.load().get("model", "")
                if current_model:
                    return [{"id": current_model}]
                # 提供常见本地模型名作为回退
                return [{"id": "gpt-3.5-turbo"}, {"id": "qwen"}, {"id": "local-model"}]
            return []
        except Exception as e:
            error_msg = str(e)
            if hasattr(e, 'response') and e.response:
                try:
                    error_data = e.response.json()
                    error_msg = error_data.get('error', {}).get('message', str(e))
                except:
                    pass
            # 本地模式下 /v1/models 不可用时，提供回退模型列表
            if is_loopback_url(base_url):
                current_model = self.config_manager.load().get("model", "")
                fallback = [{"id": current_model}] if current_model else []
                if not fallback:
                    fallback = [{"id": "gpt-3.5-turbo"}, {"id": "qwen"}, {"id": "local-model"}]
                return fallback
            return {"error": error_msg}

    def upload_files(self, files_data: list) -> list:
        result = []
        cfg = self.config_manager.load()
        vision_enabled = cfg.get("vision_enabled", False)
        for f in files_data:
            try:
                raw = base64.b64decode(f["data"])
                ext = (f.get("type") or "").lower()
                name = f.get("name", "")

                # ── 视觉文件（图片/视频）：交给视觉 API 处理 ──
                if is_visual_ext(ext):
                    if not vision_enabled:
                        result.append({
                            "name": name,
                            "size": f["size"],
                            "type": ext,
                            "error": translate_error(
                                "未配置视觉处理能力。请在设置中开启「视觉 API」并配置服务地址。"
                            ),
                        })
                        continue
                    try:
                        description = process_visual(raw, ext, cfg)
                        result.append({
                            "name": name,
                            "size": f["size"],
                            "type": ext,
                            "content": f"[视觉描述] {description}",
                        })
                    except VisionNotConfigured as ve:
                        result.append({
                            "name": name,
                            "size": f["size"],
                            "type": ext,
                            "error": translate_error(ve),
                        })
                    except Exception as ve:
                        result.append({
                            "name": name,
                            "size": f["size"],
                            "type": ext,
                            "error": translate_error(ve),
                        })
                    continue

                # ── 文本 / 文档文件：走文本提取 ──
                temp_dir = Path(self.app_dir) / "temp"
                temp_dir.mkdir(exist_ok=True)
                temp_path = temp_dir / name
                with open(temp_path, "wb") as out:
                    out.write(raw)
                try:
                    text = extract_text_from_file(str(temp_path))
                finally:
                    try:
                        os.remove(temp_path)
                    except Exception:
                        pass
                result.append({
                    "name": name,
                    "size": f["size"],
                    "type": ext,
                    "content": text
                })
            except Exception as e:
                result.append({
                    "name": f.get("name", ""),
                    "error": translate_error(e)
                })
        return result

    # ═══════════════════════════════════════════════════════════════
    #  History & Memory (session-scoped)
    # ═══════════════════════════════════════════════════════════════

    def get_initial_messages(self, session_id: str = "") -> list:
        session = self._get_session(session_id)
        return session.current_messages.copy()

    def get_memory_content(self, session_id: str = "") -> str:
        session = self._get_session(session_id)
        return session.get_memory_content(self.config_manager)

    def clear_memory(self, session_id: str = "") -> bool:
        session = self._get_session(session_id)
        return session.clear_memory()

    # ═══════════════════════════════════════════════════════════════
    #  Token usage (session-scoped)
    # ═══════════════════════════════════════════════════════════════

    def get_last_usage(self, session_id: str = "") -> dict:
        session = self._get_session(session_id)
        if not session.loop:
            return {}
        return session.loop.get_last_usage()

    def get_total_usage(self, session_id: str = "") -> dict:
        session = self._get_session(session_id)
        if not session.loop:
            return {}
        return session.loop.get_total_usage()

    # ═══════════════════════════════════════════════════════════════
    #  Plugin Management API (unchanged — global scope)
    # ═══════════════════════════════════════════════════════════════

    def get_plugins(self) -> list:
        return self.plugin_manager.get_all_plugins()

    def get_plugin_dirs(self) -> list:
        return self.plugin_manager.plugin_dirs

    def _resolve_plugin_dirs(self, dirs: list) -> list:
        """Normalise and deduplicate plugin directories by realpath.
        Filters out non-existent directories and stale paths from old sessions.

        Official plugins are no longer auto-loaded — users must add them
        explicitly via the plugin UI if desired.
        """
        resolved: List[str] = []

        for d in dirs:
            if not d or not isinstance(d, str):
                continue
            # 解析相对路径（相对于 app_dir）
            if not os.path.isabs(d):
                d = os.path.join(self.app_dir, d)
            d = os.path.normpath(d)
            # 跳过不存在的目录（例如上次会话遗留的路径）
            if not os.path.isdir(d):
                continue
            d_real = os.path.realpath(d)
            # 跳过重复目录
            if any(os.path.realpath(r) == d_real for r in resolved):
                continue
            resolved.append(d)

        return resolved

    def add_plugin_dir(self, path: str) -> str:
        cfg = self.config_manager.load()
        dirs = cfg.get("plugin_dirs", [])
        if path not in dirs:
            dirs.append(path)
            cfg["plugin_dirs"] = dirs
            self.config_manager.save(cfg)
            self.plugin_manager.set_plugin_dirs(self._resolve_plugin_dirs(dirs))
        return "ok"

    def remove_plugin_dir(self, path: str) -> str:
        cfg = self.config_manager.load()
        dirs = cfg.get("plugin_dirs", [])
        if path in dirs:
            dirs.remove(path)
            cfg["plugin_dirs"] = dirs
            self.config_manager.save(cfg)
            self.plugin_manager.set_plugin_dirs(self._resolve_plugin_dirs(dirs))
        return "ok"

    def reload_plugins(self) -> str:
        cfg = self.config_manager.load()
        dirs = cfg.get("plugin_dirs", [])
        self.plugin_manager.set_plugin_dirs(self._resolve_plugin_dirs(dirs))
        return "ok"

    def pick_plugin_dir(self) -> str:
        import webview
        try:
            result = webview.windows[0].create_file_dialog(
                webview.FileDialog.FOLDER,
                directory=self.config_manager.load().get("project_root", "")
            )
            if result and len(result) > 0:
                return result[0]
            return ""
        except Exception:
            return ""

    def get_plugin_audit_results(self) -> dict:
        return self.plugin_manager.get_audit_results()

    def get_plugin_security_config(self) -> dict:
        cfg = self.config_manager.load()
        return {
            # ★ 插件总开关（已从设置面板迁移至插件控制面板）
            "plugins_enabled": cfg.get("plugins_enabled", True),
            "audit": cfg.get("plugin_security_audit", "block"),
            "import_restrict": cfg.get("plugin_security_import_restrict", "strict"),
            "require_permissions": cfg.get("plugin_security_require_permissions", True),
            "resource_limit": cfg.get("plugin_security_resource_limit", False),
            # ★ P0-1 进程隔离
            "isolation": cfg.get("plugin_isolation", "process"),
            # ★ P0-5 签名校验
            "signature_verify": cfg.get("plugin_signature_verify", True),
            "trusted_keys": cfg.get("plugin_trusted_keys", []),
            # ★ P0-4 网络策略
            "network_policy": cfg.get("plugin_network_policy", "deny"),
            "network_url_allowlist": cfg.get("plugin_network_url_allowlist", []),
            "network_domain_allowlist": cfg.get("plugin_network_domain_allowlist", []),
            # ★ P0-8 插件工具调用审批（插件控制面板）：仅作用于插件提供的工具；
            #   原生内置工具的确认在「设置 → 原生工具确认」中配置。
            "approval_enabled": cfg.get("approval_enabled", True),
        }

    def set_plugin_security_config(self, audit: str = "block",
                                   import_restrict: str = "strict",
                                   require_permissions: bool = True,
                                   resource_limit: bool = False,
                                   isolation: str = "process",
                                   signature_verify: bool = True,
                                   trusted_keys: Optional[List[str]] = None,
                                   network_policy: str = "deny",
                                   network_url_allowlist: Optional[List[str]] = None,
                                   network_domain_allowlist: Optional[List[str]] = None,
                                   approval_enabled: bool = True,
                                   plugins_enabled: bool = True) -> str:
        cfg = self.config_manager.load()
        cfg["plugins_enabled"] = plugins_enabled
        cfg["plugin_security_audit"] = audit
        cfg["plugin_security_import_restrict"] = import_restrict
        cfg["plugin_security_require_permissions"] = require_permissions
        cfg["plugin_security_resource_limit"] = resource_limit
        cfg["plugin_isolation"] = isolation
        cfg["plugin_signature_verify"] = signature_verify
        if trusted_keys is not None:
            cfg["plugin_trusted_keys"] = trusted_keys
        cfg["plugin_network_policy"] = network_policy
        if network_url_allowlist is not None:
            cfg["plugin_network_url_allowlist"] = network_url_allowlist
        if network_domain_allowlist is not None:
            cfg["plugin_network_domain_allowlist"] = network_domain_allowlist
        cfg["approval_enabled"] = approval_enabled
        self.config_manager.save(cfg)
        self.plugin_manager.update_security_config(cfg)
        # ★ 插件总开关：启用则加载配置目录，禁用则清空所有插件
        if plugins_enabled:
            self.plugin_manager.set_plugin_dirs(self._resolve_plugin_dirs(cfg.get("plugin_dirs", [])))
        else:
            self.plugin_manager.set_plugin_dirs([])
        return "ok"

    # ═══════════════════════════════════════════════════════════════
    #  新异步架构：统计信息 API
    # ═══════════════════════════════════════════════════════════════

    def get_sandbox_pool_stats(self) -> dict:
        """获取沙箱池状态。"""
        from sandbox_pool import get_sandbox_pool
        pool = get_sandbox_pool()
        return pool.get_stats()

    def get_file_io_stats(self) -> dict:
        """获取文件 I/O 队列状态。"""
        from file_io_queue import get_file_io_queue
        queue = get_file_io_queue()
        stats = queue.get_stats()
        # 附加当前活跃文件的访问者信息
        stats["active_files"] = queue.get_all_active_files()
        return stats

    def get_lifecycle_stats(self) -> dict:
        """获取生命周期管理器状态。"""
        from lifecycle_manager import get_lifecycle_manager
        lm = get_lifecycle_manager()
        return lm.get_stats()

    def get_resource_stats(self) -> dict:
        """获取资源隔离器状态。"""
        from resource_isolator import get_resource_isolator
        isolator = get_resource_isolator()
        return isolator.get_stats()

    # ═══════════════════════════════════════════════════════════════
    #  NORP 安全系统 API
    # ═══════════════════════════════════════════════════════════════

    def get_norp_safe_logs(self, limit: int = 50) -> list:
        """获取 NORP 安全拦截日志（NORPsafe.json）。"""
        from norp_safe import get_norp_safe
        nsp = get_norp_safe(self.app_dir)
        return nsp.get_logs(limit)

    def get_norp_safe_stats(self) -> dict:
        """获取 NORP 安全系统统计。"""
        from norp_safe import get_norp_safe
        nsp = get_norp_safe(self.app_dir)
        return nsp.get_stats()

    def get_norp_safe_config(self) -> dict:
        """获取 NORP 安全系统配置状态。"""
        from norp_safe import is_norp_safe_enabled
        cfg = self.config_manager.load()
        return {
            "enabled": cfg.get("norp_safe_enabled", True),
            "runtime_enabled": is_norp_safe_enabled(),
        }

    def set_norp_safe_enabled(self, enabled: bool) -> str:
        """启用/禁用 NORP 安全系统。

        ⚠️ 禁用后所有安全检查（危险命令、UAC提权、路径越界）全部放行。
        仅建议在受信任的隔离环境中使用。前端需做二次确认。
        """
        from norp_safe import set_norp_safe_enabled
        cfg = self.config_manager.load()
        cfg["norp_safe_enabled"] = enabled
        self.config_manager.save(cfg)
        set_norp_safe_enabled(enabled)
        return "ok"

    def test_norp_safe(self, command: str) -> dict:
        """测试 NORP 拦截 — 传入命令文本，返回拦截结果。

        前端用此 API 在设置面板中测试安全系统是否正常工作。
        返回格式：{"blocked": bool, "reason": str, "threat_level": str,
                    "command": str, "norp_enabled": bool}
        """
        from norp_safe import get_norp_safe, is_norp_safe_enabled

        nsp = get_norp_safe(self.app_dir)
        result = nsp.check_command_full(command, task_id="__test__")

        return {
            "blocked": result.blocked,
            "reason": result.reason if result.blocked else "",
            "threat_level": result.threat_level.value if result.blocked else "",
            "command": command,
            "norp_enabled": is_norp_safe_enabled(),
        }

    # ═══════════════════════════════════════════════════════════════
    #  运行时完整性检测 API
    # ═══════════════════════════════════════════════════════════════

    def get_runtime_health(self) -> dict:
        """获取运行时完整性检测报告。"""
        from runtime_check import get_cached_report, run_startup_check
        report = get_cached_report()
        if report is None:
            # 如果缓存为空（不太可能），重新运行
            source_dir = os.path.dirname(os.path.abspath(__file__))
            report = run_startup_check(source_dir)
        return report.to_dict()

    # ═══════════════════════════════════════════════════════════════
    #  调试面板 API
    # ═══════════════════════════════════════════════════════════════

    def get_debug_data(self) -> dict:
        """获取 Agent 调试数据（供前端「调试」面板展示）。

        返回最近一次任务的完整调试信息：
        - react_steps: ReAct 循环时间线（思考-行动-观察）
        - tool_calls: 工具调用详情（入参/出参/耗时/沙箱路径映射）
        - security_events: NORP 安全拦截日志
        - hook_events: 插件钩子触发记录（L1-L4）
        - snapshot: 性能与状态快照（token/沙箱池/文件IO/事件队列）
        """
        from debug_logger import get_debug_logger
        dl = get_debug_logger(self.app_dir)
        return dl.get_debug_data()

    def open_debug_log_dir(self) -> str:
        """在系统文件管理器中打开调试日志目录（app_dir）。"""
        try:
            import platform
            import subprocess
            d = self.app_dir or os.getcwd()
            if not os.path.isdir(d):
                os.makedirs(d, exist_ok=True)
            if platform.system() == "Windows":
                os.startfile(d)
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", d])
            else:
                subprocess.Popen(["xdg-open", d])
            return "ok"
        except Exception as e:
            return f"error:{e}"
